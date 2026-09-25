"""Give each asset reading its own time of day, so no two readings are the same row.

The Northbridge building workbooks carry twelve readings per asset, per reading type, per day,
and every one of them is stamped at midnight: '2026-09-22T00:00:00'. Where the instrument read
the same value twice that day - a vibration sensor steady at 7.53 mm/s - the rows are identical,
and the migration's preprocessing removes identical rows as duplicates. On 24 Sep 2026 that was
232 of Harbour Point's 1,056 readings, gone before the write, reported only as "dedup ratio
78.0% < 0.80".

They are not duplicates. They are twelve readings a day with the time lost. This puts it back:
the readings of one (asset, reading type, day) are spread evenly across the day in the order the
file lists them - twelve a day is one every two hours - and nothing else about any row changes.
A group whose timestamps are already distinct is left as it is.

    python db/tools/spread_asset_readings.py                   # both building workbooks
    python db/tools/spread_asset_readings.py <path.xlsx> ...   # only these

Each file is backed up beside itself first. Rebuild the -complete workbooks afterwards with
build_complete_workbook.py, and check them with verify_workbook.py.
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import shutil
import sys

import openpyxl

DOWNLOADS = os.environ.get("NORTHBRIDGE_DIR", r"C:\Users\balap\Downloads")
DEFAULT_FILES = [
    "northbridge_B-101_harbour_point.xlsx",
    "northbridge_B-102_ashgrove_court.xlsx",
]
SHEET = "Asset_Readings"


def _parse(value) -> dt.datetime | None:
    if isinstance(value, dt.datetime):
        return value
    try:
        return dt.datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None


def spread(path: str) -> tuple[int, int]:
    """Rewrite `path` in place. Returns (readings, readings re-timed)."""
    wb = openpyxl.load_workbook(path)
    if SHEET not in wb.sheetnames:
        print(f"  {os.path.basename(path)}: no {SHEET} sheet - nothing to do")
        return 0, 0
    ws = wb[SHEET]
    header = [c.value for c in ws[1]]
    try:
        i_asset, i_type, i_at = (header.index(h) + 1 for h in
                                 ("asset_code", "reading_type", "recorded_at"))
    except ValueError as exc:
        raise SystemExit(f"  {os.path.basename(path)}: {SHEET} lacks a column: {exc}")

    groups: dict[tuple, list[tuple[int, dt.datetime, bool]]] = {}
    for row in range(2, ws.max_row + 1):
        when = _parse(ws.cell(row, i_at).value)
        if when is None:
            continue
        key = (ws.cell(row, i_asset).value, ws.cell(row, i_type).value, when.date())
        was_text = isinstance(ws.cell(row, i_at).value, str)
        groups.setdefault(key, []).append((row, when, was_text))

    total = sum(len(g) for g in groups.values())
    moved = 0
    for (_asset, _type, day), rows in groups.items():
        if len({w for _, w, _ in rows}) == len(rows):
            continue  # already distinct
        step = dt.timedelta(days=1) / len(rows)
        start = dt.datetime.combine(day, dt.time(0, 0), tzinfo=rows[0][1].tzinfo)
        for n, (row, _old, was_text) in enumerate(rows):
            new = start + step * n
            ws.cell(row, i_at).value = (new.replace(microsecond=0).isoformat() if was_text
                                        else new.replace(tzinfo=None))
            moved += 1

    if moved:
        backup = (os.path.splitext(path)[0] + ".before-spread-"
                  + dt.datetime.now().strftime("%Y%m%d-%H%M%S") + ".xlsx")
        shutil.copy2(path, backup)
        wb.save(path)
        print(f"  {os.path.basename(path)}: {moved:,} of {total:,} readings given their time "
              f"of day  (backup {os.path.basename(backup)})")
    else:
        print(f"  {os.path.basename(path)}: {total:,} readings, every timestamp already distinct")
    return total, moved


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbooks", nargs="*",
                    default=[os.path.join(DOWNLOADS, f) for f in DEFAULT_FILES])
    args = ap.parse_args()
    for p in args.workbooks:
        if not os.path.exists(p):
            print(f"  missing: {p}", file=sys.stderr)
            continue
        spread(p)


if __name__ == "__main__":
    main()
