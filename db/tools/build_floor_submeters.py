"""Build a floor-level sub-meter workbook beside a building's workbook.

A building's workbook meters it at the incoming supply: one electricity meter, one gas meter,
a year of half-hours on each. That answers "how much does the building use" and nothing about
where. This writes the companion file — `<workbook>-floorlevel_submeter.xlsx` — that puts a
sub-meter on every floor for both fuels, so the Energy page can read the building as
building → floor → meter the way the Assets page reads it as building → section → asset.

Three sheets, in the order the writer needs them:

  Building_Sections   one section per floor, named as the floor is ("Basement", "Ground",
                      "Level 1" …), carrying floor_name so the writer links floors.floor_id
  Energy_Meters       two sub-meters per floor (electricity, gas), is_sub_meter true,
                      section_name = the floor, tariff and carbon factor as the main meter's
  Meter_Readings      every half-hour of the trailing window for every sub-meter

The readings are derived from the main meters' readings in the building workbook, half-hour
by half-hour, so the floors reconcile with the incoming supply: each floor takes a fixed share
(plant-heavy in the basement, even across tenant floors) with a little noise, and the floors
together account for about 93% of electricity and 95% of gas — the remainder is lifts, risers
and common parts that a landlord does not sub-meter, and the Energy page reports that
coverage rather than pretending the floors are the building. One tenant floor carries a
deliberate night-time drift over the last five weeks so the anomaly scan has a floor-level
finding to make.

The window is 90 days by default (--days). A year per floor meter is 17,520 rows; twelve floors
and two fuels make that 420,000 rows in one sheet, which ingests slowly and says nothing more
than 90 days do about where the energy goes. The EUI is still computed from the main meters.

Reads the building workbook only. Writes nothing to any database.
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import random
import sys

import openpyxl

DEFAULT_WORKBOOKS = [
    r"C:\Users\balap\Downloads\northbridge_B-101_harbour_point.xlsx",
    r"C:\Users\balap\Downloads\northbridge_B-102_ashgrove_court.xlsx",
]
SUFFIX = "-floorlevel_submeter"

SECTIONS_HEADER = ["building_code", "name", "section_type", "floor_name", "gross_area_m2",
                   "reference_eui_kwh_m2", "reference_source"]
METERS_HEADER = ["meter_ref", "building_code", "site_ref", "meter_type", "mpan", "mprn",
                 "is_sub_meter", "section_name", "floor_name", "tariff_gbp_per_kwh",
                 "carbon_kg_per_kwh", "active", "description"]
READINGS_HEADER = ["meter_ref", "building_code", "meter_type", "reading_at", "consumption_kwh",
                   "period_minutes", "source"]

#: CIBSE TM46-style references per section type, kWh/m²/yr (electricity + fossil).
REFERENCE_EUI = {"plant": 180, "common": 205, "office": 180, "residential": 120}

#: What the floors account for of the incoming supply between them. The rest is lifts, risers,
#: external lighting and common parts a landlord does not sub-meter.
COVERAGE = {"electricity": 0.93, "gas": 0.95}
#: Basement and ground take a fixed slice; the tenant floors share what remains evenly.
FIXED_SHARE = {
    "electricity": {"Basement": 0.14, "Ground": 0.10},
    "gas": {"Basement": 0.62, "Ground": 0.08},
}
NOISE = 0.05
#: A tenant floor whose night-time electricity creeps up over the last five weeks, so the
#: scan has something to find at floor level. Applied to the highest tenant floor.
DRIFT_WEEKS = 5
DRIFT_NIGHT_UPLIFT = 0.25


def floor_names(n: int) -> list[str]:
    """Basement, Ground, Level 1 … the way plenum_cafm.floors names them for these buildings."""
    if n <= 0:
        return []
    names = ["Basement", "Ground"]
    names += [f"Level {i}" for i in range(1, max(0, n - 2) + 1)]
    return names[:n] if n >= 2 else names[:1]


def floor_code(name: str) -> str:
    if name == "Basement":
        return "B"
    if name == "Ground":
        return "G"
    return "L" + name.split()[-1].zfill(2)


def section_type_for(name: str, primary_use: str) -> str:
    if name == "Basement":
        return "plant"
    if name == "Ground":
        return "common"
    return "residential" if (primary_use or "").lower().startswith("resid") else "office"


def read_sheet(wb, name: str) -> tuple[list[str], list[list]]:
    ws = wb[name]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    return header, [list(r) for r in rows]


def shares(fuel: str, names: list[str]) -> dict[str, float]:
    fixed = FIXED_SHARE[fuel]
    tenants = [n for n in names if n not in fixed]
    remaining = COVERAGE[fuel] - sum(fixed.get(n, 0.0) for n in names if n in fixed)
    out = {n: fixed[n] for n in names if n in fixed}
    for n in tenants:
        out[n] = remaining / len(tenants) if tenants else 0.0
    return out


def _parse_at(v) -> dt.datetime:
    if isinstance(v, dt.datetime):
        return v if v.tzinfo else v.replace(tzinfo=dt.timezone.utc)
    s = str(v).replace("Z", "+00:00")
    d = dt.datetime.fromisoformat(s)
    return d if d.tzinfo else d.replace(tzinfo=dt.timezone.utc)


def build(path: str, *, days: int, seed: int) -> str:
    rng = random.Random(seed)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    bh, brows = read_sheet(wb, "Buildings")
    b = dict(zip(bh, brows[0]))
    code, name = str(b["building_code"]), str(b["name"])
    site_ref = b.get("site_ref")
    n_floors = int(float(b["floors"]))
    gia_m2 = float(b["gross_internal_area_m2"] or 0) or float(b["gross_area_sqft"]) / 10.7639
    use = str(b.get("primary_use") or "Commercial")
    names = floor_names(n_floors)
    area_each = round(gia_m2 / len(names), 1)

    mh, mrows = read_sheet(wb, "Energy_Meters")
    mains = {}
    for r in mrows:
        m = dict(zip(mh, r))
        if str(m.get("is_sub_meter", "false")).lower() in ("true", "1", "yes"):
            continue
        mains[str(m["meter_type"]).lower()] = m
    if not mains:
        raise SystemExit(f"{os.path.basename(path)}: Energy_Meters has no main meter")

    rh, rrows = read_sheet(wb, "Meter_Readings")
    ri = {h: i for i, h in enumerate(rh)}
    latest = max(_parse_at(r[ri["reading_at"]]) for r in rrows)
    window_start = latest - dt.timedelta(days=days)
    drift_start = latest - dt.timedelta(weeks=DRIFT_WEEKS)
    by_fuel: dict[str, list[tuple[dt.datetime, float, int]]] = {}
    for r in rrows:
        at = _parse_at(r[ri["reading_at"]])
        if at < window_start:
            continue
        by_fuel.setdefault(str(r[ri["meter_type"]]).lower(), []).append(
            (at, float(r[ri["consumption_kwh"]] or 0), int(r[ri["period_minutes"]] or 30)))
    for fuel in by_fuel:
        by_fuel[fuel].sort()

    # Sheets ---------------------------------------------------------------------------
    sections = []
    for fname in names:
        st = section_type_for(fname, use)
        sections.append([code, fname, st, fname, area_each, REFERENCE_EUI[st], f"CIBSE TM46 {st}"])

    meters, readings = [], []
    drift_floor = next((n for n in reversed(names) if n.startswith("Level")), None)
    for fuel, main in mains.items():
        if fuel not in by_fuel:
            print(f"  {name}: no {fuel} readings in the last {days} days — no {fuel} sub-meters")
            continue
        main_ref = str(main.get("meter_ref") or main.get("mpan") or main.get("mprn"))
        prefix = main_ref.rsplit("-", 1)[0]          # NB-B-101-E0 -> NB-B-101
        letter = "E" if fuel == "electricity" else "G"
        share = shares(fuel, names)
        for fname in names:
            ref = f"{prefix}-{letter}-{floor_code(fname)}"
            meters.append([
                ref, code, site_ref, fuel,
                ref if fuel == "electricity" else None, ref if fuel == "gas" else None,
                "true", fname, fname,
                main.get("tariff_gbp_per_kwh"), main.get("carbon_kg_per_kwh"), "true",
                f"{name} {fname} {fuel} sub-meter",
            ])
            w = share[fname]
            for at, kwh, minutes in by_fuel[fuel]:
                v = kwh * w * (1.0 + rng.uniform(-NOISE, NOISE))
                if fuel == "electricity" and fname == drift_floor and at >= drift_start \
                        and (at.hour >= 22 or at.hour < 6):
                    v *= 1.0 + DRIFT_NIGHT_UPLIFT
                readings.append([ref, code, fuel, at.strftime("%Y-%m-%dT%H:%M:%SZ"),
                                 round(v, 3), minutes, "bms"])

    out_path = os.path.splitext(path)[0] + SUFFIX + ".xlsx"
    try:
        with open(out_path, "ab"):
            pass
    except PermissionError:
        out_path = os.path.splitext(path)[0] + SUFFIX + "-" + dt.datetime.now().strftime("%Y%m%d-%H%M") + ".xlsx"
        print(f"  {name}: target is open elsewhere - writing {os.path.basename(out_path)}")

    out = openpyxl.Workbook(write_only=True)
    for title, header, rows in (("Building_Sections", SECTIONS_HEADER, sections),
                                ("Energy_Meters", METERS_HEADER, meters),
                                ("Meter_Readings", READINGS_HEADER, readings)):
        ws = out.create_sheet(title)
        ws.append(header)
        for r in rows:
            ws.append(r)
    out.save(out_path)

    print(f"  {name} ({code}): {len(names)} floors, {len(meters)} sub-meters, "
          f"{len(readings):,} readings over {days} days ({window_start.date()} to {latest.date()})")
    print(f"     floors account for {int(COVERAGE['electricity']*100)}% of electricity, "
          f"{int(COVERAGE['gas']*100)}% of gas; night drift on {drift_floor} electricity "
          f"from {drift_start.date()}")
    print(f"     -> {out_path}")
    return out_path


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbooks", nargs="*", default=DEFAULT_WORKBOOKS)
    ap.add_argument("--days", type=int, default=90, help="trailing window of half-hours per sub-meter")
    ap.add_argument("--seed", type=int, default=101)
    args = ap.parse_args()
    for i, p in enumerate(args.workbooks):
        if not os.path.exists(p):
            print(f"  missing: {p}", file=sys.stderr)
            continue
        build(p, days=args.days, seed=args.seed + i)


if __name__ == "__main__":
    main()
