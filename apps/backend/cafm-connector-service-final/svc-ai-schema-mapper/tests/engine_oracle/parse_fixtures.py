"""The files the parse oracle reads: each one a case the Python parse handles some particular way.

``write_fixtures(dir)`` writes every fixture and returns their paths. Workbooks are made with
openpyxl (what the platform itself writes) and, where openpyxl cannot say it, edited XML.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import PatternFill


def _book(path: Path, sheets: dict, *, epoch=None, setup=None) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if epoch is not None:
        wb.epoch = epoch
    for title, rows in sheets.items():
        ws = wb.create_sheet(title=title)
        for row in rows:
            ws.append(row)
    if setup:
        setup(wb)
    wb.save(path)
    return path


def _formats(path: Path, header: list[str], values: list, formats: list[str]) -> Path:
    """One column per number format: the same kind of value, shown through each format."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Formats"
    ws.append(header)
    for r, v in enumerate(values, start=2):
        for c, fmt in enumerate(formats, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.number_format = fmt
    wb.save(path)
    return path


# ── hand-built workbooks: shared strings and a date style, each part replaceable ──
_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
_DECL = '<?xml version="1.0" encoding="{enc}" standalone="yes"?>\n'
_STRINGS = ["code", "name", "when"] + [f"A{i}" for i in range(1, 6)] + [f"Pump {i}" for i in range(1, 6)]


def _sst(strings: list, enc: str = "UTF-8") -> str:
    body = "".join(f"<si><t>{s}</t></si>" for s in strings)
    return f'{_DECL.format(enc=enc)}<sst xmlns="{_NS}" count="{len(strings)}" uniqueCount="{len(strings)}">{body}</sst>'


def _sheet(enc: str = "UTF-8", cell_b2: str = '<c r="B2" t="s"><v>8</v></c>') -> str:
    rows = ['<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c><c r="C1" t="s"><v>2</v></c></row>']
    for i in range(1, 6):
        b = cell_b2 if i == 1 else f'<c r="B{i + 1}" t="s"><v>{7 + i}</v></c>'
        rows.append(f'<row r="{i + 1}"><c r="A{i + 1}" t="s"><v>{2 + i}</v></c>{b}'
                    f'<c r="C{i + 1}" s="1"><v>{45657 + i}.5</v></c></row>')
    return f'{_DECL.format(enc=enc)}<worksheet xmlns="{_NS}"><sheetData>{"".join(rows)}</sheetData></worksheet>'


_CT = (_DECL.format(enc="UTF-8") + '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
       '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
       '<Default Extension="xml" ContentType="application/xml"/>'
       '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
       '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
       '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
       '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/></Types>')
_RELS = (_DECL.format(enc="UTF-8") + f'<Relationships xmlns="{_PKG}"><Relationship Id="rId1" '
         f'Type="{_REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
_WB = (_DECL.format(enc="UTF-8") + f'<workbook xmlns="{_NS}" xmlns:r="{_REL}"><sheets>'
       '<sheet name="Data" sheetId="1" r:id="rId1"/></sheets></workbook>')
_WB_RELS = (_DECL.format(enc="UTF-8") + f'<Relationships xmlns="{_PKG}">'
            f'<Relationship Id="rId1" Type="{_REL}/worksheet" Target="worksheets/sheet1.xml"/>'
            f'<Relationship Id="rId2" Type="{_REL}/sharedStrings" Target="sharedStrings.xml"/>'
            f'<Relationship Id="rId3" Type="{_REL}/styles" Target="styles.xml"/></Relationships>')
_STYLES = (_DECL.format(enc="UTF-8") + f'<styleSheet xmlns="{_NS}"><fonts count="1"><font><sz val="11"/></font></fonts>'
           '<fills count="1"><fill><patternFill patternType="none"/></fill></fills><borders count="1"><border/></borders>'
           '<cellStyleXfs count="1"><xf numFmtId="0"/></cellStyleXfs><cellXfs count="2"><xf numFmtId="0" xfId="0"/>'
           '<xf numFmtId="22" xfId="0" applyNumberFormat="1"/></cellXfs></styleSheet>')


def _parts_book(path: Path, **parts) -> Path:
    """A workbook from its parts; any part given as bytes replaces the intact one as is."""
    import zipfile

    base = {"[Content_Types].xml": _CT, "_rels/.rels": _RELS, "xl/workbook.xml": _WB,
            "xl/_rels/workbook.xml.rels": _WB_RELS, "xl/worksheets/sheet1.xml": _sheet(),
            "xl/sharedStrings.xml": _sst(_STRINGS), "xl/styles.xml": _STYLES}
    base.update({k.replace("__", "/"): v for k, v in parts.items()})
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in base.items():
            z.writestr(name, data if isinstance(data, bytes) else data.encode("utf-8"))
    return path


def _damaged_and_encoded(d: Path) -> list[Path]:
    """Exports cut short or saved in another encoding. Python refuses each damaged_* file (the
    part the cut lands in decides the words); it reads the tolerated_* and encoded_* ones."""
    s, t, y = _sheet().encode(), _sst(_STRINGS).encode(), _STYLES.encode()
    wb, rels = _WB.encode(), _WB_RELS.encode()
    row4 = s.index(b'<row r="4"')
    latin = list(_STRINGS)
    latin[8] = "Café München"
    sheet = "xl__worksheets__sheet1.xml"
    return [
        _parts_book(d / "damaged_cut_in_attribute.xlsx", **{sheet: s[: row4 + 25]}),
        _parts_book(d / "damaged_cut_in_tag.xlsx", **{sheet: s[: row4 + 4]}),
        _parts_book(d / "damaged_cut_in_value.xlsx", **{sheet: s[: row4 + s[row4:].index(b"<v>") + 4]}),
        _parts_book(d / "damaged_cut_after_cell.xlsx", **{sheet: s[: row4 + s[row4:].index(b"</c>") + 4]}),
        _parts_book(d / "damaged_cut_between_rows.xlsx", **{sheet: s[:row4]}),
        _parts_book(d / "damaged_cut_in_comment.xlsx", **{sheet: s[:row4] + b"<!-- a comment that never"}),
        _parts_book(d / "damaged_shared_strings.xlsx", **{"xl__sharedStrings.xml": t[: t.index(b"<si><t>Pump 2")]}),
        _parts_book(d / "damaged_styles.xlsx", **{"xl__styles.xml": y[: y.index(b"</cellXfs>")]}),
        _parts_book(d / "damaged_styles_after_cellxfs.xlsx",
                    **{"xl__styles.xml": y[: y.index(b"</cellXfs>") + len(b"</cellXfs>")]}),
        _parts_book(d / "damaged_workbook_part.xlsx", **{"xl__workbook.xml": wb[: wb.index(b"</sheets>") + 9]}),
        _parts_book(d / "damaged_relationships.xlsx",
                    **{"xl___rels__workbook.xml.rels": rels[: rels.index(b'<Relationship Id="rId2"')]}),
        _parts_book(d / "damaged_utf16_sheet.xlsx", **{sheet: _sheet("UTF-16").encode("utf-16")}),
        _parts_book(d / "damaged_utf16_shared_strings.xlsx",
                    **{"xl__sharedStrings.xml": _sst(_STRINGS).replace(_DECL.format(enc="UTF-8"), "").encode("utf-16")}),
        _parts_book(d / "damaged_invalid_utf8_shared_strings.xlsx",
                    **{"xl__sharedStrings.xml": t.replace(b"Pump 1", b"Pump \xe9 1")}),
        _parts_book(d / "damaged_invalid_utf8_sheet.xlsx",
                    **{sheet: _sheet(cell_b2='<c r="B2" t="inlineStr"><is><t>Zxrich</t></is></c>').encode()
                       .replace(b"Zxrich", b"Z\xfcrich")}),
        _parts_book(d / "tolerated_cut_after_sheetdata.xlsx", **{sheet: s[: s.index(b"</sheetData>") + 12]}),
        _parts_book(d / "tolerated_cut_after_shared_strings.xlsx", **{"xl__sharedStrings.xml": t[: t.index(b"</sst>") + 6]}),
        _parts_book(d / "encoded_latin1_shared_strings.xlsx",
                    **{"xl__sharedStrings.xml": _sst(latin, "ISO-8859-1").encode("latin-1")}),
        _parts_book(d / "encoded_cp1252_shared_strings.xlsx",
                    **{"xl__sharedStrings.xml": _sst(latin, "windows-1252").encode("cp1252")}),
        _parts_book(d / "encoded_latin1_sheet.xlsx",
                    **{sheet: _sheet("ISO-8859-1", '<c r="B2" t="inlineStr"><is><t>Zürich</t></is></c>').encode("latin-1")}),
    ]


def write_fixtures(d: Path) -> list[Path]:
    d.mkdir(parents=True, exist_ok=True)
    out: list[Path] = []
    add = out.append

    # ── workbooks ──
    add(_book(d / "banner_rows.xlsx", {"Assets": [["Asset register — exported 1 Oct"], [],
                                                  ["code", "name", "site"], ["A1", "Pump", "B-101"],
                                                  ["A2", "Fan", "B-102"]]}))

    def _merge(wb):
        wb["Report"].merge_cells("A1:C1")
    add(_book(d / "merged_title.xlsx", {"Report": [["Quarterly report"], ["id", "value", "unit"],
                                                   ["1", 10, "kWh"], ["2", 20.5, "kWh"]]}, setup=_merge))
    add(_book(d / "dup_headers.xlsx", {"Codes": [["Code", "Code", "code", "Name"], ["X1", "X1", "x1", "a"],
                                                 ["X2", "X2", "x2", "b"], ["X3", "X3", "x3", "c"]]}))
    add(_book(d / "empty_headers.xlsx", {"Gaps": [["id", None, "name", None], ["1", "a", "x", "q"],
                                                  ["2", "b", "y", None]]}))
    add(_book(d / "numeric_headers.xlsx", {"Years": [[2024, 1.5, dt.datetime(2024, 1, 31), "label"],
                                                     [1, 2, 3, "a"], [4, 5, 6, "b"]]}))
    builtin = ["mm-dd-yy", "d-mmm-yy", "d-mmm", "mmm-yy", "h:mm AM/PM", "h:mm:ss AM/PM", "h:mm", "h:mm:ss",
               "m/d/yy h:mm", "mm:ss", "[h]:mm:ss", "mmss.0"]
    add(_formats(d / "dates_builtin.xlsx", [f"f{i}" for i in range(len(builtin))],
                 [dt.datetime(2024, 2, 29, 13, 45, 30), dt.datetime(1900, 3, 1, 0, 0), 45000.75, 0.5, 1.25],
                 builtin))
    custom = ["dd/mm/yyyy hh:mm", "yyyy-mm-dd", "mmm-yy", "[h]:mm", "h:mm AM/PM", '"Date: "dd/mm',
              "[$-409]d-mmm-yyyy", "0.00", "#,##0", "0%", "@", '[Red]0.0;[Blue]-0.0', "dddd", "ss.000"]
    add(_formats(d / "dates_custom_formats.xlsx", [f"c{i}" for i in range(len(custom))],
                 [dt.datetime(2026, 10, 1, 9, 5, 7), 45566.0, 1.5, 0.0001, 12345.678], custom))
    add(_book(d / "dates_1904.xlsx", {"Mac": [["when", "n"], [dt.datetime(2024, 5, 6, 7, 8, 9), 1],
                                              [dt.date(1904, 1, 2), 2]]}, epoch=openpyxl.utils.datetime.CALENDAR_MAC_1904))
    add(_book(d / "times_and_durations.xlsx", {"Times": [["t", "d"], [dt.time(10, 30), dt.timedelta(hours=1, minutes=30)],
                                                         [dt.time(23, 59, 59, 500000), dt.timedelta(days=2, seconds=5)],
                                                         [dt.time(0, 0), dt.timedelta(0)]]}))
    add(_book(d / "booleans_errors.xlsx", {"Flags": [["b", "e", "n"], [True, "#N/A", 1], [False, "#DIV/0!", 2],
                                                     [True, "#VALUE!", 3]]}))
    add(_book(d / "formulas_no_cache.xlsx", {"Calc": [["a", "b", "sum"], [1, 2, "=A2+B2"], [3, 4, "=A3+B3"]]}))

    def _rich(wb):
        ws = wb["Rich"]
        ws["A2"] = CellRichText(["plain ", TextBlock(InlineFont(b=True), "bold"), " end"])
    add(_book(d / "rich_text.xlsx", {"Rich": [["text", "n"], [None, 1], ["simple", 2]]}, setup=_rich))

    def _hide(wb):
        wb["Secret"].sheet_state = "hidden"
    add(_book(d / "hidden_sheet.xlsx", {"Visible": [["a"], ["1"]], "Secret": [["b"], ["2"]]}, setup=_hide))
    add(_book(d / "post_write_sheets.xlsx", {"Vendors": [["vendor_code", "vendor_name"], ["V1", "Acme"]],
                                             "ContractTerms": [["k", "v"], ["sla", "4h"]],
                                             "README": [["note"], ["read me"]],
                                             "Invoices": [["inv"], ["I-1"]]}))
    add(_book(d / "numbers.xlsx", {"Nums": [["n"], [0.1], [1e-7], [1e16], [123456789012], [-0.0], [2 ** 53 + 1],
                                            ["007"], [1.0], [3.14159], [-12.5], [1e300], [123.456e-10]]}))

    def _na_text(wb):
        c = wb["NA"]["E2"]
        c._value, c.data_type = "#N/A", "s"  # the text "#N/A", not the error
    add(_book(d / "na_tokens.xlsx", {"NA": [["a", "b", "c", "d", "e", "f"], ["NA", "N/A", "null", "None", "x", "   "],
                                            ["nan", "NaN", "-", "n/a", "<NA>", ""]]}, setup=_na_text))
    add(_book(d / "empty_and_header_only.xlsx", {"Empty": [], "HeaderOnly": [["a", "b"]], "Data": [["x"], ["1"]]}))

    def _styled(wb):
        ws = wb["Wide"]
        fill = PatternFill("solid", fgColor="FFFF00")
        for col in range(4, 27):
            for row in range(1, 6):
                ws.cell(row=row, column=col).fill = fill
    add(_book(d / "wide_styled_empty_columns.xlsx", {"Wide": [["a", "b", "c"], ["1", "2", "3"], ["4", "5", "6"]]},
              setup=_styled))
    add(_book(d / "unicode.xlsx", {"Ünïcode ✓": [["naïve", "日本語", "emoji"], ["é", "テスト", "🏗️"],
                                                 [" nbsp ", "ß", "𝔘𝔫𝔦"]]}))

    out += _damaged_and_encoded(d)

    # ── delimited text ──
    def text(name: str, content: str, encoding: str = "utf-8"):
        p = d / name
        p.write_bytes(content.encode(encoding))
        add(p)

    text("comma.csv", "code,name,qty\nA1,Pump,3\nA2,Fan,\n")
    text("semicolon.csv", "code;name;qty\nA1;Pump;3\nA2;Fan;4\n")
    text("tab.tsv", "code\tname\tnote\nA1\tPump\tred, large\nA2\tFan\tsmall\n")
    text("pipe.csv", "code|name\nA1|Pump\nA2|Fan\n")
    text("quoted_newlines.csv", 'code,note\nA1,"line one\nline two"\nA2,"say ""hi"""\nA3,"a,b"\n')
    p = d / "bom_crlf.csv"
    p.write_bytes(b"\xef\xbb\xbfcode,name\r\nA1,Pump\r\nA2,Fan\r\n")
    add(p)
    text("blank_and_whitespace_lines.csv", "code,name\nA1,Pump\n\n   \nA2,Fan\n \t \n  A3, Spaced \n")
    text("na_tokens.csv", "a,b,c,d,e,f,g\nNA,N/A,null,None,#N/A,,nan\n-,NULL,<NA>,n/a,-NaN,  ,x\n")
    text("short_rows.csv", "a,b,c\n1,2,3\n4,5\n6\n")
    text("too_many_fields.csv", "a,b\n1,2\n3,4,5\n")
    text("trailing_comma_data_rows.csv", "a,b\n1,2,\n3,4,\n")
    text("dup_and_empty_headers.csv", "id,,id,name,,Unnamed: 1\n1,a,2,x,b,c\n")
    text("cp1252.csv", "name,price\nCafé,€5\nIt’s,€7\n", encoding="cp1252")
    text("latin1.csv", "name,city\nJosé,München\nRenée,Zürich\n", encoding="latin-1")
    return out
