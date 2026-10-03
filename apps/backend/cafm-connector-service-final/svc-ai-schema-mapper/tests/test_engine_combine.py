"""A multi-file upload's combined workbook, made by `hoist-engine combine`, reads as Python's does.

Python's combine (multi_upload.combine_in_python, today's start-with-upload-multi code) and the
engine's are run on the same uploads; each combined workbook is then read as the run's ingest reads
it (header row, dtype=str, sanitised names, records), and the two must agree sheet by sheet, row by
row — and fail alike, with the same words, where Python fails.
"""
from __future__ import annotations

import asyncio
import io

import openpyxl
import pandas as pd
import pytest

from src import multi_upload as mu
from src.engine import client, store
from src.excel_parser import ExcelWorkbook
from src.graph.nodes.ingest_node import _sanitize_column_names, _sanitize_records

MIGRATION = "6c1d2e3f-4a5b-4c6d-8e7f-9a0b1c2d3e4f"
needs_engine = pytest.mark.skipif(not client.engine_available(),
                                  reason="hoist-engine not built (engine/scripts/dev.sh build-linux)")


def workbook(sheets: dict) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def read_back(data: bytes) -> dict:
    wb = ExcelWorkbook(io.BytesIO(data))
    try:
        out = {}
        for name in wb.sheet_names:
            df = _sanitize_column_names(wb.read(name, header=wb.header_row(name), dtype=str))
            out[name] = (list(df.columns), _sanitize_records(df.to_dict(orient="records")))
        return out
    finally:
        wb.close()


UPLOADS = [
    ("b101 export.xlsx", workbook({
        "Sites": [["Site Report"], [None], ["code", " name ", "code"], ["S1", "Harbour Point", "S1"], ["S2", None, "x"]],
        "Assets": [["asset_code", "installed", "qty"], ["A-1", "2024-01-15", 3], ["A-2", None, 2.5],
                   ["=formula", "#REF!", "NA"]],
        "Contract Terms": [["vendor", "sla"], ["V1", "4h"]],
    })),
    ("readings;semi.csv", b"meter;date;kwh\nM1;2025-04-01;1.2\nM2;2025-04-01;\n"),
    ("tabbed.tsv", b"a\tb\tb\n1\t2\t3\n\t\t\n"),
    ("Sheet1 only.xlsx", workbook({"Sheet1": [["x", "y"], ["1", "2"]]})),
    ("again.csv", b"code,name\nS9,Quay\n"),
    ("Sites.csv", b"only\nv\n"),
]


@pytest.fixture()
def engine_dir(monkeypatch, tmp_path):
    monkeypatch.setattr(store, "ENGINE_DIR", tmp_path / "engine")


@needs_engine
def test_the_engine_combines_the_uploads_as_python_does(engine_dir):
    py_bytes, py_names = mu.combine_in_python(UPLOADS)
    got = asyncio.run(mu.combine_with_engine(UPLOADS, MIGRATION))
    assert got is not None
    go_bytes, go_names = got
    assert go_names == py_names
    assert read_back(go_bytes) == read_back(py_bytes)
    assert py_names == ["Sites", "Assets", "Contract Terms", "readings;semi", "tabbed", "Sheet1 only", "again", "Sites_2"]


@needs_engine
@pytest.mark.parametrize("uploads", [
    [("good.csv", b"a\n1\n"), ("bad.xlsx", b"PK\x03\x04 not a workbook"), ("worse.csv", b"a,b\n1,2,3\n")],
    [("broken.csv", b'a,b\n"x,1\n'), ("ok.xlsx", workbook({"S": [["a"], ["1"]]}))],
    [("ctl.csv", b"a\nbell \x07 here\n")],
], ids=["first-failing-file", "pandas-cannot-read", "a-value-excel-cannot-hold"])
def test_the_engine_fails_where_python_fails_with_the_same_words(engine_dir, uploads):
    with pytest.raises(Exception) as py_err:
        mu.combine_in_python(uploads)
    with pytest.raises(mu.CombineError) as go_err:
        asyncio.run(mu.combine_with_engine(uploads, MIGRATION))
    assert str(go_err.value) == str(py_err.value)


def test_only_csv_tsv_and_real_workbooks_go_to_the_engine(monkeypatch):
    monkeypatch.setattr(mu, "_engine_ready", lambda: True)
    assert mu.engine_can_combine([("a.csv", b"x"), ("b.XLSX", b"PK\x03\x04..."), ("c.tsv", b"y"), ("d.xlsm", b"PK\x03\x04")])
    assert not mu.engine_can_combine([("a.csv", b"x"), ("old.xls", b"\xd0\xcf\x11\xe0")])
    assert not mu.engine_can_combine([("locked.xlsx", b"\xd0\xcf\x11\xe0")])  # an encrypted workbook is OLE, not zip
    monkeypatch.setattr(mu, "_engine_ready", lambda: False)
    assert not mu.engine_can_combine([("a.csv", b"x")])


def test_an_engine_that_crashes_leaves_the_combine_to_python(engine_dir, monkeypatch):
    async def crash(*a, **k):
        raise client.EngineError("internal", "combine exited 2: boom")

    monkeypatch.setattr(mu, "run_engine", crash)
    assert asyncio.run(mu.combine_with_engine([("a.csv", b"a\n1\n")], MIGRATION)) is None
