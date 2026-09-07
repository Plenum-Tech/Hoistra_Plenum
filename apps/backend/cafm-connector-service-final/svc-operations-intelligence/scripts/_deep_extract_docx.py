"""Deep-extract UK pack docx cards including nested cell text."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

DOCX = Path(r"C:\Users\balap\Documents\UK_Compliance_Certification_Pack_v1.1.docx")
OUT = Path("scripts/_uk_pack_docx_deep.txt")


def iter_block_items(parent):
    from docx.oxml.ns import qn

    parent_elm = parent.element.body if hasattr(parent, "element") else parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def cell_full_text(cell) -> str:
    parts: list[str] = []
    for p in cell.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    for t in cell.tables:
        for row in t.rows:
            for c in row.cells:
                ct = cell_full_text(c)
                if ct:
                    parts.append(ct)
    return "\n".join(parts)


def main() -> None:
    doc = Document(str(DOCX))
    lines: list[str] = []
    scope = None
    trade = None
    card_n = 0
    for table in doc.tables:
        t0 = table.rows[0].cells[0].text.strip()
        if "SECTION 1" in t0.upper():
            scope = "Building"
            lines.append(f"\n##### {t0}")
            continue
        if "SECTION 2" in t0.upper():
            scope = "Vendor"
            lines.append(f"\n##### {t0}")
            continue
        if len(table.rows) == 1:
            trade = t0
            lines.append(f"\n### {trade}")
            continue
        if not scope:
            continue
        card_n += 1
        title = t0
        body = cell_full_text(table.rows[1].cells[0]) if len(table.rows) > 1 else ""
        lines.append(f"\n==== CARD {card_n} [{scope}] {title} ====")
        lines.append(body[:4000] if body else "(empty body)")

    OUT.write_text("\n".join(lines), encoding="utf-8")
    print("cards", card_n, "wrote", OUT)


if __name__ == "__main__":
    main()
