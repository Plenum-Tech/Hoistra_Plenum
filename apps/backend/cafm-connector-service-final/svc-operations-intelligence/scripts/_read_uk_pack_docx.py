from pathlib import Path

p = Path(r"C:\Users\balap\Documents\UK_Compliance_Certification_Pack_v1.1.docx")
print("exists", p.exists(), "size", p.stat().st_size if p.exists() else 0)

try:
    from docx import Document
except ImportError:
    import subprocess
    import sys

    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    from docx import Document

doc = Document(str(p))
print("=== PARAS", len(doc.paragraphs), "===")
for i, para in enumerate(doc.paragraphs):
    t = para.text.strip()
    if t:
        style = para.style.name if para.style else ""
        print(f"P{i}|{style}|{t}")

print("=== TABLES", len(doc.tables), "===")
for ti, table in enumerate(doc.tables):
    print(f"--- TABLE {ti} rows={len(table.rows)} cols={len(table.columns)} ---")
    for ri, row in enumerate(table.rows):
        cells = [c.text.strip().replace("\n", " | ") for c in row.cells]
        print(f"T{ti}R{ri}|" + " || ".join(cells))
