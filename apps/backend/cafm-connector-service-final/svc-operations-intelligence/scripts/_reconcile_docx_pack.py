"""Reconcile Word UK pack docx vs shipped JSON pack."""
from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document

DOCX = Path(r"C:\Users\balap\Documents\UK_Compliance_Certification_Pack_v1.1.docx")
JSON = Path("src/reference/uk_compliance_pack_v1_1.json")

SECTION_HEADERS = {
    "SECTION 1",
    "SECTION 2",
    "HOW TO READ",
}
TRADE_HEADERS = re.compile(r"^[A-K]\.\s+")


def extract_cards(doc: Document) -> list[dict]:
    cards: list[dict] = []
    scope = None
    trade = None
    for table in doc.tables:
        if len(table.rows) < 1:
            continue
        title = table.rows[0].cells[0].text.strip()
        if not title:
            continue
        up = title.upper()
        if "SECTION 1" in up or "BUILDING-LEVEL" in up:
            scope = "Building"
            continue
        if "SECTION 2" in up or "VENDOR" in up:
            scope = "Vendor"
            continue
        if "HOW TO READ" in up:
            continue
        if TRADE_HEADERS.match(title) or (
            len(table.rows) == 1 and title.endswith(("SAFETY", "ELECTRICAL", "GAS SAFETY", "LEGIONELLA", "SYSTEMS", "LOLER)", "ASBESTOS", "ENVIRONMENTAL", "(F-GAS)", "DOCUMENTATION", "storeys", "CONTRACTORS", "TRANSPORTATION CONTRACTORS", "LEGIONELLA CONTRACTORS", "(F-GAS)", "CONSTRUCTION CONTRACTORS", "SECURITY CONTRACTORS", "PEST CONTROL"))
        ):
            # trade section header tables are 1-row
            if len(table.rows) == 1:
                trade = title
                continue
        # certificate card: usually 2 rows, title in row0
        if len(table.rows) >= 1 and scope and not TRADE_HEADERS.match(title):
            # skip pure section headers already handled
            if title.startswith("SECTION"):
                continue
            if len(table.rows) == 1 and title.isupper() or TRADE_HEADERS.match(title):
                trade = title
                continue
            body = ""
            if len(table.rows) > 1:
                body = table.rows[1].cells[0].text.strip()
            # Heuristic: cards have body content or known certificate-ish titles
            if body or any(
                k in title
                for k in (
                    "Certificate",
                    "Report",
                    "Register",
                    "Assessment",
                    "Record",
                    "Policy",
                    "Survey",
                    "Card",
                    "Licence",
                    "License",
                    "Accreditation",
                    "Competency",
                    "Membership",
                    "Insurance",
                    "Evidence",
                    "Authorisation",
                    "Pack",
                )
            ):
                cards.append(
                    {
                        "scope": scope,
                        "trade": trade,
                        "title": title,
                        "body_len": len(body),
                    }
                )
    return cards


def main() -> None:
    doc = Document(str(DOCX))
    cards = extract_cards(doc)
    pack = json.loads(JSON.read_text(encoding="utf-8"))
    jtypes = pack["certificate_types"]

    print("DOCX cards:", len(cards))
    b = [c for c in cards if c["scope"] == "Building"]
    v = [c for c in cards if c["scope"] == "Vendor"]
    print("  Building:", len(b))
    print("  Vendor:", len(v))
    print("JSON types:", len(jtypes), f"({pack['type_counts']})")
    print()
    print("=== DOCX BUILDING TITLES ===")
    for c in b:
        print(" -", c["title"])
    print("=== DOCX VENDOR TITLES ===")
    for c in v:
        print(" -", c["title"])

    # Also dump raw 2-row table titles with section tracking (simpler)
    print("\n=== SIMPLE WALK ===")
    scope = None
    titles = []
    for table in doc.tables:
        t0 = table.rows[0].cells[0].text.strip()
        if "SECTION 1" in t0.upper():
            scope = "Building"
            continue
        if "SECTION 2" in t0.upper():
            scope = "Vendor"
            continue
        if len(table.rows) == 1:
            continue  # headers
        if scope and t0:
            titles.append((scope, t0))
    print("2-row cards:", len(titles))
    from collections import Counter

    print(Counter(s for s, _ in titles))
    for s, t in titles:
        print(f"{s[0]}|{t}")


if __name__ == "__main__":
    main()
