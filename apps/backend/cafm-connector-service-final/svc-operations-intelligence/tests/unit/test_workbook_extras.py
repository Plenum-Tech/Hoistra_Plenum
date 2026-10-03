"""A single end-to-end workbook's contract terms and invoices, read for the engines.

The migration sets Contract_Terms and Invoice_Lines aside; workbook_extras reads them from the
stored upload and hands them to the contract ingest and the invoice matcher.
"""
import io

import openpyxl

from src.engines.contract_performance import workbook_extras as wx


def _book() -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "Vendors"
    wb.active.append(["vendor_code", "vendor_name"])
    ct = wb.create_sheet("Contract_Terms")
    ct.append(["contract_ref", "contract_name", "vendor_code", "vendor_name", "signed_date", "term", "value", "source", "clause", "page"])
    for term, value, src, clause in [
            ("P1 response", "4 hours", "contract", "5.2(a)"), ("P2 response", "1 business day", "contract", "5.2(b)"),
            ("Labour rate — standard", "£68 / hr", "contract", "Sch. 2"), ("Labour rate — out of hours", "£102 / hr", "contract", "Sch. 2"),
            ("Completion target", "95%", "contract", "5.4"), ("Parts mark-up cap", "15%", "default", None)]:
        ct.append(["AM-2024-HVAC-07", "Mechanical PPM · Bishopsgate", "APXM", "Apex Mechanical", "2024-04-27", term, value, src, clause, 11])
    il = wb.create_sheet("Invoice Lines")          # any spelling of the sheet name
    il.append(["invoice_no", "vendor_code", "wo_code", "labour_hours", "labour_rate", "amount"])
    il.append(["INV-8841", "APXM", "WO-B-301-4262", 32, 85, 3808])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_only_the_two_sheets_are_read_whatever_their_spelling():
    sheets = wx.read_sheets(_book())
    assert set(sheets) == {"contractterms", "invoicelines"}
    assert len(sheets["contractterms"]) == 6
    assert sheets["invoicelines"][0]["wo_code"] == "WO-B-301-4262"


def test_terms_become_the_extractors_shape_and_defaults_stay_out():
    ext = wx.extraction_from_terms(wx.read_sheets(_book())["contractterms"])
    assert ext["sla_response_p1_hours"] == 4
    assert ext["sla_response_p2_hours"] == 24          # a business day is 24 elapsed hours
    assert ext["labour_hour_rate"] == 68 and ext["overtime_rate"] == 102
    assert ext["kpi_clauses_json"]["Completion target"] == {"value": "95%", "clause": "5.4", "page": 11}
    assert "parts_pricing_json" not in ext             # a default term is left to the engine


def test_a_workbook_without_them_reads_empty():
    wb = openpyxl.Workbook()
    out = io.BytesIO()
    wb.save(out)
    assert wx.read_sheets(out.getvalue()) == {}


def test_plant_telemetry_sheets_are_read_and_their_values_parse():
    from datetime import date, datetime, timezone
    from src.engines.energy import workbook_telemetry as wt

    wb = openpyxl.Workbook()
    wb.active.title = "Chiller_Readings"
    wb.active.append(["asset_code", "reading_at", "kw_input", "cooling_load_rt"])
    wb.active.append(["B-301-CHILLER-101", "2026-09-14T15:00:00Z", 212.4, 280.0])
    bms = wb.create_sheet("BMS Trends")
    bms.append(["building_code", "zone", "recorded_at", "heating_pct", "cooling_pct"])
    bms.append(["B-301", "L20 North", "2026-09-20T01:15:00Z", 38.0, 31.5])
    wd = wb.create_sheet("Weather_Degree_Days")
    wd.append(["building_code", "month", "hdd", "cdd"])
    wd.append(["B-301", "2026-09-01", 45.2, 60.4])
    out = io.BytesIO()
    wb.save(out)
    sheets = wx.read_sheets(out.getvalue())
    assert set(sheets) == {"chillerreadings", "bmstrends", "weatherdegreedays"}
    assert wt._at(sheets["chillerreadings"][0]["reading_at"]) == datetime(2026, 9, 14, 15, tzinfo=timezone.utc)
    assert wt._month("2026-09-17") == wt._month("2026-09") == date(2026, 9, 1)
    assert wt._at("not a time") is None and wt._f("") is None


def test_a_finished_read_is_kept_as_counts():
    from src.engines.contract_performance.workbook_extras_runner import summarise

    s = summarise({"found": True, "contracts": [{}, {}],
                   "invoices": [{"matched": 9, "held": 1}, {"matched": 3, "held": 0}],
                   "skipped": [{"reason": "already verified"}, {"reason": "already verified"}],
                   "telemetry": {"chiller_readings": 960, "bms_samples": 8070, "skipped": []}})
    assert s["contracts"] == 2 and s["invoices"] == 2
    assert s["lines_held"] == 1 and s["lines_matched"] == 12
    assert s["skipped"] == 2 and s["skip_reasons"] == ["already verified"]
    assert s["telemetry"] == {"chiller_readings": 960, "bms_samples": 8070}
