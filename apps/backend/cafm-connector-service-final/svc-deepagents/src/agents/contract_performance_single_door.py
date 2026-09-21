"""
Single Door → Feature B (Contract Performance) auto-route.

When PM uploads FM contracts, framework agreements, POs, or invoices via the
Orchestrator, classify and call svc-operations-intelligence extract APIs
(Relationships Agent mid-layer) so structured Contract/Invoice fields are
drafted for HITL confirmation.
"""
from __future__ import annotations

import base64
import re
from pathlib import Path
from typing import Any, Literal

import structlog

from ..config import settings
from ..http_client import request as _request

log = structlog.get_logger(__name__)

ContractPerfKind = Literal["contract", "invoice"]

_CONTRACT_NAME = re.compile(
    r"\b(contract|framework|agreement|purchase[\s_-]?order|\bpo[\s_-]|"
    r"sla|framework[\s_-]?agreement|msa|sow)\b",
    re.I,
)
_INVOICE_NAME = re.compile(
    r"\b(invoice|inv[\s_-]|bill|credit[\s_-]?note|remittance)\b",
    re.I,
)
_CONTRACT_MSG = re.compile(
    r"\b(contract|framework agreement|purchase order|\bpo\b|sla param|"
    r"vendor agreement|fm contract)\b",
    re.I,
)
_INVOICE_MSG = re.compile(
    r"\b(invoice|bill|verify invoice|invoice check|invoice match)\b",
    re.I,
)


def classify_contract_performance_doc(
    file_path: str,
    user_query: str | None = None,
) -> ContractPerfKind | None:
    """
    Detect Feature B uploads from filename + user message.
    Invoice wins over contract if both match (invoice often cites PO/contract).
    """
    # Normalize separators so word-boundary regex works on Framework_Agreement etc.
    name = re.sub(r"[_\-.]+", " ", Path(file_path).name)
    # ...and split CamelCase, because the keywords below are word-bounded and procurement
    # portals export names like "ProcurementContractData" with the word glued inside a
    # compound. "Contract" is right there in the filename, but contract cannot see it,
    # so the document silently misses Feature B entirely unless the user happens to say
    # "contract" in their message.
    # Digits close the boundary just as tightly as case does — "Agreement2024" hides the
    # word from agreement — so break letter/digit runs apart as well.
    name = re.sub(r"(?<=[a-z0-9])(?=[A-Z])|(?<=[A-Za-z])(?=[0-9])|(?<=[0-9])(?=[A-Za-z])", " ", name)
    msg = (user_query or "").strip()

    inv_hit = bool(_INVOICE_NAME.search(name) or (msg and _INVOICE_MSG.search(msg)))
    con_hit = bool(_CONTRACT_NAME.search(name) or (msg and _CONTRACT_MSG.search(msg)))

    if inv_hit and not con_hit:
        return "invoice"
    if con_hit and not inv_hit:
        return "contract"
    if inv_hit and con_hit:
        if _INVOICE_NAME.search(name) or (msg and _INVOICE_MSG.search(msg)):
            return "invoice"
        return "contract"
    return None


# A 40-page cap silently truncated the UKRI FM contract at page 40 of 52 — dropping the
# KPI Schedule entirely. "Service Credit" appears 14 times in that document and 0 times in
# its first 40 pages. FM contracts put their SLA annexes at the back, so a cap that reads
# only the front discards precisely the part the performance engine exists to read.
_MAX_PDF_PAGES = 250


def _pdf_base64(file_path: str, *, limit: int = 28_000_000) -> str | None:
    """Base64 of the PDF, so the extractor can read tables and scanned pages visually.

    Returns None for non-PDFs and for files past the API's document limit — in both cases
    extraction falls back to the text it was already given rather than failing.
    """
    path = Path(file_path)
    if path.suffix.lower() != ".pdf":
        return None
    try:
        raw = path.read_bytes()
    except OSError as exc:  # noqa: BLE001
        log.warning("single_door.cp.pdf_bytes_failed", error=str(exc)[:200])
        return None
    encoded = base64.b64encode(raw).decode()
    if len(encoded) >= limit:
        log.warning(
            "single_door.cp.pdf_too_large_for_vision",
            name=path.name,
            encoded_len=len(encoded),
        )
        return None
    return encoded


def extract_text_from_upload(file_path: str, *, max_chars: int = 50_000) -> str:
    """Best-effort text for Relationships Agent extraction."""
    path = Path(file_path)
    ext = path.suffix.lower()
    try:
        if ext in {".txt", ".md", ".csv"}:
            return path.read_text(encoding="utf-8", errors="replace")[:max_chars]
        if ext in {".docx"}:
            try:
                from docx import Document  # type: ignore

                doc = Document(str(path))
                parts = [p.text for p in doc.paragraphs if p.text]
                for table in doc.tables:
                    for row in table.rows:
                        parts.append(" | ".join(c.text for c in row.cells))
                return "\n".join(parts)[:max_chars]
            except Exception as exc:  # noqa: BLE001
                log.warning("single_door.cp.docx_read_failed", error=str(exc)[:200])
        if ext == ".pdf":
            try:
                from pypdf import PdfReader  # type: ignore

                reader = PdfReader(str(path))
                parts = []
                for page in reader.pages[:_MAX_PDF_PAGES]:
                    parts.append(page.extract_text() or "")
                text = "\n".join(parts).strip()
                if text:
                    return text[:max_chars]
            except Exception as exc:  # noqa: BLE001
                log.warning("single_door.cp.pdf_read_failed", error=str(exc)[:200])
            # Second rung, for PDFs pypdf reads as empty (many scanned or
            # producer-quirky certificates). This used to import pdfplumber, which is not a
            # dependency of this service — so the rung always raised ModuleNotFoundError and
            # the document fell through to the filename stub with no body text at all. That
            # is what left a generically-named certificate unclassifiable. PyMuPDF is
            # installed and reads these reliably.
            try:
                import pymupdf  # type: ignore

                with pymupdf.open(str(path)) as pdf:
                    parts = [page.get_text() or "" for page in pdf.pages(0, _MAX_PDF_PAGES)]
                text = "\n".join(parts).strip()
                if text:
                    return text[:max_chars]
                log.warning("single_door.cp.pdf_text_empty", path=path.name)
            except Exception as exc:  # noqa: BLE001
                log.warning("single_door.cp.pymupdf_failed", error=str(exc)[:200])
    except OSError as exc:
        log.warning("single_door.cp.read_failed", error=str(exc)[:200])

    # Fallback: filename + stub so extract still runs heuristics / Claude on thin context
    return f"Document filename: {path.name}\n(User uploaded via Single Door Orchestrator.)"


#: Entries in `field_sources` that are not contract TERMS. `building_link` records whether
#: the document's building could be matched against the register — an outcome, not a term —
#: and counting it would inflate the denominator the panel prints.
_NON_TERM_SOURCES = frozenset({"building_link"})


def term_counts(body: dict | None) -> tuple[int | None, int | None]:
    """(terms read from the document, terms in total) — the panel's own arithmetic.

    The reply used to count keys in the model's raw output, which includes `vendor_name` and
    `signed_date`. Neither is a contract term the Vendors panel lists, so one ingest produced
    two numbers: "extracted 14 parameters" in the chat and "12 of 17 terms were read" on the
    panel, with nothing to tell the reader which was the real coverage.

    `field_sources` is the record both surfaces should speak from — the engine writes it, the
    panel already reads it, and "contract" in it is exactly what "the document said this"
    means. Returns (None, None) when the response carries no field_sources: an older engine or
    a partial body is a reason to say nothing, not to guess.
    """
    sources = (body or {}).get("field_sources")
    if not isinstance(sources, dict) or not sources:
        return None, None
    terms = {k: v for k, v in sources.items() if k not in _NON_TERM_SOURCES}
    read = sum(1 for v in terms.values() if str(v).strip().lower() == "contract")
    return read, len(terms)


def ingest_summary(
    *,
    name: str,
    read: int | None,
    total: int | None,
    defaults: int,
    vendor_note: str,
    params_id: str | None = None,
) -> str:
    """What the chat says after a contract ingest, in the panel's numbers.

    A document that yielded NOTHING is called out separately. Moreland's ingest read 0 of 16
    terms — correctly, because the file is a property management agreement and not a service
    contract — and the reply was worded identically to an ingest with a few gaps. The reader
    confirmed 16 platform defaults as agreed values on the strength of it.
    """
    if read is None or total is None:
        head = f"Contract Performance B1: parameters extracted from {name} — {defaults} defaulted."
    elif read == 0:
        head = (
            f"Contract Performance B1: NO CONTRACT TERMS were read from {name}. Every one of "
            f"the {total} terms fell back to a platform default, so this may not be a service "
            f"contract — confirming it would make {total} assumed values binding on the vendor."
        )
    else:
        head = (
            f"Contract Performance B1: read {read} of {total} terms from {name} "
            f"— {defaults} not stated by the contract and defaulted."
        )
    return f"{head}{vendor_note} Scoring stays blocked until these are confirmed."


def vendor_note_for(vendor: dict | None) -> str:
    """What the reply says about the supplier the contract names — for EVERY outcome.

    This used to cover three statuses and fall silent on the rest. On 17 Sep 2026 the
    backend emitted `could_not_create` (it was squashing the company id to an integer before
    the INSERT), the contract was written with no vendor, and the reply read "Ingestion
    complete" with nothing about the supplier at all. The reader's first sight of the
    problem was a Vendors page with their contract missing from it.

    A contract with no vendor has no card to appear under, so that is the fact to state.
    Any status this function does not recognise is treated as that failure rather than as
    nothing to say: a new status added upstream and not handled here must fail loud.
    """
    v = vendor or {}
    status = str(v.get("status") or "")
    name = str(v.get("vendor_name") or "").strip()
    if status == "supplied":
        return ""
    if status == "created":
        return f" Vendor '{name}' was new and is now registered."
    if status == "matched":
        return f" Linked to existing vendor '{name}'."
    if status == "not_named_in_contract" or not name:
        return " No service provider named in the document — contract is unlinked."
    if status == "legacy_id_unlinkable":
        # The vendor is right there in the register. Saying "could not be registered" would
        # send the reader off to create a duplicate; the actual blocker is that half the
        # register predates uuid ids and no column that points at a vendor can hold one.
        legacy = str(v.get("legacy_vendor_id") or "").strip()
        return (
            f" '{name}' IS already in the register (id {legacy or 'a legacy code'}), but that "
            f"id predates uuids and the contract's vendor field cannot hold it, so this "
            f"contract is saved WITHOUT a vendor. Do not re-register the supplier — the row "
            f"needs its id migrated."
        )
    return (
        f" The document names '{name}' but that supplier could not be registered, so the "
        f"contract is saved WITHOUT a vendor and will not appear under Vendors until one is "
        f"attached."
    )


async def route_contract_performance_upload(
    *,
    file_path: str,
    organization_id: str | None = None,
    user_query: str | None = None,
    document_id: str | None = None,
) -> dict[str, Any] | None:
    """
    If the upload is a contract/PO/invoice, call Feature B extract APIs.
    Returns None when not a Feature B document.
    """
    kind = classify_contract_performance_doc(file_path, user_query)
    if kind is None:
        return None

    # Whole document, not the first 50k: extraction downstream selects the relevant
    # passages itself, and can only select from what it is given.
    source_text = extract_text_from_upload(file_path, max_chars=400_000)
    base = settings.operations_intelligence_base_url.rstrip("/")
    name = Path(file_path).name

    try:
        if kind == "contract":
            resp = await _request(
                "POST",
                base,
                "/api/contract-performance/contracts/extract",
                service="operations_intelligence",
                timeout=90.0,
                json={
                    "source_text": source_text,
                    "pdf_base64": _pdf_base64(file_path),
                    "organization_id": organization_id,
                    "document_id": document_id,
                    "contract_ref": Path(file_path).stem,
                    "auto_ingest": True,
                },
            )
            body = resp.json()
            params_id = body.get("parameters_id") or (body.get("ingest") or {}).get(
                "parameters", {}
            ).get("id")
            defaults = list(body.get("defaults_used") or [])
            vendor = body.get("vendor") or {}
            vendor_note = vendor_note_for(vendor)
            # Counted off field_sources, which is what the Vendors panel counts. Counting the
            # model's raw keys here is what made one ingest report "14" in the chat and
            # "12 of 17" on the panel.
            read, total = term_counts(body)
            summary = ingest_summary(
                name=name, read=read, total=total, defaults=len(defaults),
                vendor_note=vendor_note, params_id=params_id,
            )
            return {
                "kind": "contract",
                "ok": bool(body.get("ok", True)),
                "summary": summary,
                "result": body,
                "parameters_id": params_id,
                # Everything the chat card needs to show the extraction inline, so the PM
                # can see what was read without navigating away to find out.
                "extraction": {
                    "document_name": name,
                    "parameters_id": params_id,
                    "vendor": vendor,
                    "extracted": body.get("extracted") or {},
                    "field_confidence": body.get("field_confidence") or {},
                    "field_sources": body.get("field_sources") or {},
                    "defaults_used": defaults,
                    "preview_with_defaults": body.get("preview_with_defaults") or {},
                    "source_truncated": body.get("source_truncated"),
                },
            }

        # invoice
        resp = await _request(
            "POST",
            base,
            "/api/contract-performance/invoices/extract-verify",
            service="operations_intelligence",
            timeout=90.0,
            json={
                "source_text": source_text,
                "organization_id": organization_id,
                "document_id": document_id,
                # The filename, offered as a label of last resort and not as the invoice's
                # identity. It carries the session id and changes on every upload, so it
                # keyed the register on how a file happened to be named and made the same
                # invoice look like two. The server reads the number off the document.
                "invoice_ref_fallback": Path(file_path).stem,
            },
        )
        body = resp.json()
        verification = body.get("verification") or body
        matched = verification.get("matched_count")
        flagged = verification.get("flagged_count")
        summary = (
            f"Contract Performance B3: verified invoice {name} — "
            f"matched={matched}, flagged={flagged}. "
            f"Review flagged lines in Approvals (Feature B); £500+ went through Adversary."
        )
        return {
            "kind": "invoice",
            "ok": bool(body.get("ok", True)),
            "summary": summary,
            "result": body,
            "verification_id": verification.get("verification_id"),
        }
    except Exception as exc:  # noqa: BLE001
        log.error(
            "single_door.contract_performance.route_failed",
            kind=kind,
            file=name,
            error=str(exc)[:300],
        )
        return {
            "kind": kind,
            "ok": False,
            "summary": f"Feature B {kind} route failed for {name}: {exc}",
            "error": str(exc)[:500],
            "result": None,
        }
