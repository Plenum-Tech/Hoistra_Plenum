"""Preview and transactionally ingest a complete building workbook.

The importer is intentionally separate from the agent-facing single-record API.
This keeps a large owner-approved batch behind an exact hash confirmation while
preserving the existing narrow tool contract used by the model.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


ORGANIZATION_ID = "hw1-org-a"
SOURCE_BUILDING_ID = "BLD-CDF-0001"

# Exact headers prevent a changed workbook from silently entering a previously
# approved batch under a familiar filename.
SHEET_HEADERS: dict[str, tuple[str, ...]] = {
    "building": ("building_id", "name", "source_dataset", "source_building_id", "source_site_id", "address_line", "city", "region", "country", "latitude", "longitude", "timezone", "primary_use", "sub_use", "sector", "gross_internal_area_sqm", "gross_internal_area_sqft", "occupied_floors", "levels_total", "year_built", "occupants", "heating_type", "cooling_type", "display_energy_certificate_rating", "metered_utilities", "ownership", "occupier", "tenure_note", "fm_model", "data_window_source", "data_window_demo", "demo_time_shift_days", "source_license", "synthetic_fields_note"),
    "floors": ("floor_id", "building_id", "name", "level_index", "area_sqm", "description", "floor_type"),
    "spaces": ("space_id", "building_id", "floor_id", "name", "space_type", "area_sqm", "capacity_persons", "notes"),
    "systems": ("system_id", "building_id", "name", "description", "brick_class"),
    "assets": ("asset_id", "asset_tag", "name", "category", "brick_class", "uniclass_ss", "system_id", "space_id", "building_id", "manufacturer", "model", "serial_number", "install_date", "expected_life_years", "capacity", "capacity_unit", "criticality", "condition_grade", "status", "replacement_cost_gbp", "maintained_by_vendor_id", "parent_asset_id", "serves_space_ids", "warranty_expiry", "age_years_at_demo_end", "remaining_life_years", "data_provenance"),
    "meters": ("meter_id", "building_id", "name", "utility", "meter_role", "unit", "space_id", "parent_meter_id", "serves", "data_provenance", "supply_point_ref_synthetic", "make_model", "install_date", "amr_enabled", "interval", "reading_type"),
    "tariffs": ("tariff_id", "meter_id", "utility", "valid_from", "valid_to", "unit_rate_gbp", "non_commodity_rate_gbp", "standing_charge_gbp_per_day", "notes", "data_provenance"),
    "energy_monthly": ("meter_id", "building_id", "month", "kwh", "hours", "hours_ok", "data_completeness_pct", "utility", "blended_rate_gbp_per_kwh", "cost_gbp", "co2e_kg"),
    "energy_annual_summary": ("utility", "year_demo", "kwh", "cost_gbp", "co2e_kg", "completeness", "kwh_per_sqm", "note"),
    "meter_events": ("event_id", "meter_id", "building_id", "event_type", "start", "end", "duration_hours", "severity", "description", "observed_in_real_data", "source_start", "linked_work_order_id"),
    "meter_readings_daily_water": ("meter_id", "building_id", "date", "value", "unit", "data_provenance", "quality_flag"),
    "vendors": ("vendor_id", "trading_name", "legal_name", "vendor_type", "services", "website", "primary_discipline", "country", "relationship_note", "company_registration_number", "approved_status", "onboarded_date", "payment_terms_days", "account_manager", "account_manager_email", "account_manager_phone", "helpdesk_phone", "performance_rating_12m"),
    "contacts": ("contact_id", "name", "role", "organisation", "contact_type", "vendor_id", "email", "phone", "is_fictional_person"),
    "contracts": ("contract_id", "building_id", "vendor_id", "title", "contract_type", "start_date", "end_date", "annual_value_gbp", "scope", "sla_summary", "kpis", "service_credit_cap_pct", "renewal_notice_days", "status", "currency", "data_provenance"),
    "sla_matrix": ("priority", "definition", "response_hours", "rectification_hours"),
    "vendor_certificates": ("certificate_id", "vendor_id", "vendor_name", "certificate_type", "category", "issuer", "issuer_is_fictional", "certificate_number", "issue_date", "expiry_date", "coverage_amount_gbp", "scope", "status", "file_name", "verified_by", "verified_date", "data_provenance"),
    "ppm_schedules": ("schedule_id", "asset_category", "task", "frequency", "statutory", "standard_reference", "vendor_id", "duration_hours", "contract_id"),
    "work_orders": ("work_order_id", "building_id", "wo_type", "discipline", "title", "description", "asset_id", "space_id", "floor_id", "system_id", "vendor_id", "contract_id", "ppm_schedule_id", "priority", "status", "requested_by_contact_id", "created_at", "scheduled_date", "response_due_at", "responded_at", "rectification_due_at", "completed_at", "sla_response_met", "sla_rectification_met", "labour_hours", "labour_cost_gbp", "parts_cost_gbp", "total_cost_gbp", "currency", "chargeable", "downtime_hours", "resolution", "root_cause", "linked_meter_event_id", "linked_alarm_id", "data_provenance"),
    "work_order_lines": ("line_id", "work_order_id", "line_type", "description", "quantity", "unit", "unit_cost_gbp", "line_total_gbp", "vendor_id"),
    "bms_alarms": ("alarm_id", "building_id", "asset_id", "point_name", "alarm_text", "severity", "raised_at", "cleared_at", "acknowledged_by", "linked_work_order_id", "nuisance", "source_system"),
    "compliance_documents": ("document_id", "building_id", "document_type", "title", "statutory", "standard_reference", "vendor_id", "asset_ids", "issued_date", "next_due_date", "outcome", "status", "file_name", "data_provenance"),
    "edges": ("from_id", "from_label", "relationship", "to_id", "to_label"),
}

SHEET_KEYS: dict[str, tuple[str, ...]] = {
    "building": ("building_id",), "floors": ("floor_id",), "spaces": ("space_id",),
    "systems": ("system_id",), "assets": ("asset_id",), "meters": ("meter_id",),
    "tariffs": ("tariff_id",), "energy_monthly": ("meter_id", "month"),
    "energy_annual_summary": ("utility", "year_demo"), "meter_events": ("event_id",),
    "meter_readings_daily_water": ("meter_id", "date"), "vendors": ("vendor_id",),
    "contacts": ("contact_id",), "contracts": ("contract_id",), "sla_matrix": ("priority",),
    "vendor_certificates": ("certificate_id",), "ppm_schedules": ("schedule_id",),
    "work_orders": ("work_order_id",), "work_order_lines": ("line_id",),
    "bms_alarms": ("alarm_id",), "compliance_documents": ("document_id",),
    "edges": ("from_id", "relationship", "to_id"),
}

# Nullable references are checked only when populated. Requiring them to resolve
# catches accidental partial uploads without inventing replacement identifiers.
FOREIGN_KEYS: tuple[tuple[str, str, str, str], ...] = (
    ("floors", "building_id", "building", "building_id"),
    ("spaces", "building_id", "building", "building_id"),
    ("spaces", "floor_id", "floors", "floor_id"),
    ("systems", "building_id", "building", "building_id"),
    ("assets", "building_id", "building", "building_id"),
    ("assets", "system_id", "systems", "system_id"),
    ("assets", "space_id", "spaces", "space_id"),
    ("assets", "maintained_by_vendor_id", "vendors", "vendor_id"),
    ("assets", "parent_asset_id", "assets", "asset_id"),
    ("meters", "building_id", "building", "building_id"),
    ("meters", "space_id", "spaces", "space_id"),
    ("meters", "parent_meter_id", "meters", "meter_id"),
    ("tariffs", "meter_id", "meters", "meter_id"),
    ("energy_monthly", "meter_id", "meters", "meter_id"),
    ("energy_monthly", "building_id", "building", "building_id"),
    ("meter_events", "meter_id", "meters", "meter_id"),
    ("meter_events", "building_id", "building", "building_id"),
    ("meter_events", "linked_work_order_id", "work_orders", "work_order_id"),
    ("meter_readings_daily_water", "meter_id", "meters", "meter_id"),
    ("meter_readings_daily_water", "building_id", "building", "building_id"),
    ("contacts", "vendor_id", "vendors", "vendor_id"),
    ("contracts", "building_id", "building", "building_id"),
    ("contracts", "vendor_id", "vendors", "vendor_id"),
    ("vendor_certificates", "vendor_id", "vendors", "vendor_id"),
    ("ppm_schedules", "vendor_id", "vendors", "vendor_id"),
    ("ppm_schedules", "contract_id", "contracts", "contract_id"),
    ("work_orders", "building_id", "building", "building_id"),
    ("work_orders", "asset_id", "assets", "asset_id"),
    ("work_orders", "space_id", "spaces", "space_id"),
    ("work_orders", "floor_id", "floors", "floor_id"),
    ("work_orders", "system_id", "systems", "system_id"),
    ("work_orders", "vendor_id", "vendors", "vendor_id"),
    ("work_orders", "contract_id", "contracts", "contract_id"),
    ("work_orders", "ppm_schedule_id", "ppm_schedules", "schedule_id"),
    ("work_orders", "requested_by_contact_id", "contacts", "contact_id"),
    ("work_orders", "linked_meter_event_id", "meter_events", "event_id"),
    ("work_orders", "linked_alarm_id", "bms_alarms", "alarm_id"),
    ("work_order_lines", "work_order_id", "work_orders", "work_order_id"),
    ("work_order_lines", "vendor_id", "vendors", "vendor_id"),
    ("bms_alarms", "building_id", "building", "building_id"),
    ("bms_alarms", "asset_id", "assets", "asset_id"),
    ("bms_alarms", "linked_work_order_id", "work_orders", "work_order_id"),
    ("compliance_documents", "building_id", "building", "building_id"),
    ("compliance_documents", "vendor_id", "vendors", "vendor_id"),
)

SENSITIVE_SHEETS = {"contacts"}
CONTACT_FIELDS = {"account_manager", "account_manager_email", "account_manager_phone", "helpdesk_phone", "email", "phone"}
ALLOWED_REFERENCE_LITERALS = {
    ("ppm_schedules", "vendor_id"): {"IN-HOUSE"},
    ("work_orders", "vendor_id"): {"IN-HOUSE"},
    ("work_order_lines", "vendor_id"): {"IN-HOUSE"},
    ("work_orders", "requested_by_contact_id"): {"OCCUPANT-HELPDESK"},
}
MULTI_REFERENCE_FIELDS = {("meter_events", "linked_work_order_id")}


@dataclass(frozen=True)
class SourceRow:
    sheet: str
    row_number: int
    key: str
    fields: dict[str, Any]


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True)


def digest(value: Any) -> str:
    return hashlib.sha256(canonical(value).encode()).hexdigest()


def file_digest(path: Path) -> str:
    hasher = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            hasher.update(block)
    return hasher.hexdigest()


def row_key(sheet: str, fields: dict[str, Any]) -> str:
    parts = [str(fields.get(column) or "").strip() for column in SHEET_KEYS[sheet]]
    if any(not part for part in parts):
        raise ValueError(f"{sheet}: required identity is missing")
    return "|".join(parts)


def read_rows(path: Path) -> tuple[dict[str, list[SourceRow]], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    errors: list[str] = []
    parsed: dict[str, list[SourceRow]] = {}
    try:
        actual = tuple(workbook.sheetnames)
        expected = tuple(SHEET_HEADERS)
        if actual != expected:
            errors.append(f"sheet order/names changed: expected {expected}, received {actual}")
        for sheet, expected_headers in SHEET_HEADERS.items():
            if sheet not in workbook.sheetnames:
                continue
            worksheet = workbook[sheet]
            values = worksheet.iter_rows(values_only=True)
            headers = tuple(str(value) if value is not None else "" for value in next(values, ()))
            if headers != expected_headers:
                errors.append(f"{sheet}: headers changed")
                continue
            seen: set[str] = set()
            rows: list[SourceRow] = []
            for row_number, values_row in enumerate(values, start=2):
                fields = {name: json_value(value) for name, value in zip(headers, values_row)}
                try:
                    key = row_key(sheet, fields)
                except ValueError as exc:
                    errors.append(f"{sheet} row {row_number}: {exc}")
                    continue
                if key in seen:
                    errors.append(f"{sheet} row {row_number}: duplicate identity {key}")
                    continue
                seen.add(key)
                rows.append(SourceRow(sheet, row_number, key, fields))
            parsed[sheet] = rows
    finally:
        workbook.close()
    return parsed, errors


def validate_references(rows: dict[str, list[SourceRow]]) -> list[str]:
    errors: list[str] = []
    indexes = {
        (sheet, column): {row.fields.get(column) for row in sheet_rows}
        for sheet, sheet_rows in rows.items()
        for column in SHEET_HEADERS[sheet]
    }
    for sheet, column, target_sheet, target_column in FOREIGN_KEYS:
        targets = indexes.get((target_sheet, target_column), set())
        for row in rows.get(sheet, []):
            value = row.fields.get(column)
            if value in (None, "") or value in ALLOWED_REFERENCE_LITERALS.get((sheet, column), set()):
                continue
            values = str(value).split(";") if (sheet, column) in MULTI_REFERENCE_FIELDS else [value]
            for reference in values:
                reference = reference.strip() if isinstance(reference, str) else reference
                if reference not in targets:
                    errors.append(f"{sheet} row {row.row_number}: {column}={reference!r} has no {target_sheet}.{target_column}")

    # The edge table is a cross-domain relationship graph. Both endpoints must
    # exist somewhere in the workbook even when their entity types differ.
    identifiers = {
        value
        for sheet, sheet_rows in rows.items()
        if sheet != "edges"
        for row in sheet_rows
        for column in SHEET_KEYS[sheet]
        for value in (row.fields.get(column),)
        if value not in (None, "")
    }
    for row in rows.get("edges", []):
        for column in ("from_id", "to_id"):
            if row.fields[column] not in identifiers:
                errors.append(f"edges row {row.row_number}: {column}={row.fields[column]!r} is unresolved")
    return errors


def sanitize(fields: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in fields.items() if key not in CONTACT_FIELDS}


def source_id(workbook_sha: str, row: SourceRow) -> str:
    return "workbook-source-" + hashlib.sha256(f"{workbook_sha}|{row.sheet}|{row.key}".encode()).hexdigest()


def record_id(workbook_sha: str, row: SourceRow) -> str:
    return "workbook-record-" + hashlib.sha256(f"{workbook_sha}|{row.sheet}|{row.key}".encode()).hexdigest()


def source_payload(workbook_sha: str, row: SourceRow) -> dict[str, Any]:
    fields = row.fields
    return {
        "original": canonical(fields),
        "fields": fields,
        "content_sha256": digest(fields),
        "incoming_sha256": digest({"workbook_sha256": workbook_sha, "sheet": row.sheet, "key": row.key, "fields": fields}),
        "version": 1,
        "information_type": "workbook_" + row.sheet,
        "scope_verified": True,
        "sensitive": row.sheet in SENSITIVE_SHEETS,
        "workbook_sha256": workbook_sha,
        "sheet": row.sheet,
        "row_number": row.row_number,
        "source_building_id": SOURCE_BUILDING_ID,
    }


def record_payload(workbook_sha: str, row: SourceRow) -> dict[str, Any]:
    source = source_id(workbook_sha, row)
    return {
        "fields": row.fields,
        "source_id": source,
        "version": 1,
        "adapter_key": row.key,
        "accepted": True,
        "scope_verified": True,
        "workbook_sha256": workbook_sha,
        "sheet": row.sheet,
        "row_number": row.row_number,
    }


def read_projection(row: SourceRow, workbook_sha: str, target_building: str) -> list[dict[str, Any]]:
    fields = sanitize(row.fields)
    reference = f"workbook:{workbook_sha}:{row.sheet}:{row.row_number}"
    base = {
        "organization_id": ORGANIZATION_ID,
        "building_id": target_building,
        "content_building_ids": [target_building],
        "scope_verified": True,
        "contains_sensitive_data": False,
        "owner_user_id": None,
        "category": None,
    }

    def item(identifier: str, kind: str, payload: dict[str, Any]) -> dict[str, Any]:
        return {**base, "record_id": identifier, "kind": kind,
                "payload": {**payload, "source_reference": reference}}

    if row.sheet == "contacts":
        return []
    if row.sheet == "work_orders":
        work_order = str(fields["work_order_id"])
        return [
            item(work_order, "work_order_status", {
                "work_order_id": work_order, "status": fields.get("status"),
                "asset_id": fields.get("asset_id"), "due_at": fields.get("rectification_due_at") or fields.get("scheduled_date"),
            }),
            item("sla-" + work_order, "work_order_sla", {
                "work_order_id": work_order,
                "finding": {"priority": fields.get("priority"), "response_met": fields.get("sla_response_met"),
                            "rectification_met": fields.get("sla_rectification_met"),
                            "response_due_at": fields.get("response_due_at"),
                            "rectification_due_at": fields.get("rectification_due_at")},
            }),
        ]
    if row.sheet == "assets":
        asset = str(fields["asset_id"])
        return [
            item(asset, "asset_information", {"asset_id": asset, "model": fields.get("model"), "details": fields}),
            item("condition-" + asset, "asset_conditions", {"asset_id": asset,
                 "condition": fields.get("condition_grade"), "observed_at": fields.get("warranty_expiry")}),
        ]
    if row.sheet == "vendor_certificates":
        certificate = str(fields["certificate_id"])
        return [
            item(certificate, "certificate_status", {"certificate_id": certificate, "subject_type": "vendor",
                 "subject_id": fields.get("vendor_id"), "status": fields.get("status"), "expiry_date": fields.get("expiry_date")}),
            item("certificate-document-" + certificate, "certificate_documents", {"title": fields.get("certificate_type"),
                 "content": fields, "expiry_date": fields.get("expiry_date")}),
        ]
    if row.sheet == "compliance_documents":
        identifier = str(fields["document_id"])
        return [item(identifier, "compliance_evidence", {"title": fields.get("title"), "content": fields,
                     "certificate_id": None})]
    if row.sheet == "contracts":
        identifier = str(fields["contract_id"])
        return [item(identifier, "contract_documents", {"title": fields.get("title"), "content": fields,
                     "vendor_id": fields.get("vendor_id")})]
    if row.sheet == "vendors":
        identifier = str(fields["vendor_id"])
        return [item(identifier, "vendor_performance", {"vendor_id": identifier,
                     "finding": {"trading_name": fields.get("trading_name"),
                                 "approved_status": fields.get("approved_status"),
                                 "performance_rating_12m": fields.get("performance_rating_12m")}})]

    generic_id = f"vasiliki:{row.sheet}:{row.key}"
    return [item(generic_id, "building_records", {"record_type": row.sheet, "title": row.key, "details": fields})]


def flatten(rows: dict[str, list[SourceRow]]) -> list[SourceRow]:
    return [row for sheet in SHEET_HEADERS for row in rows.get(sheet, [])]


def run_psql(sql: str) -> str:
    result = subprocess.run(
        ["docker", "exec", "-i", "hoistra-hw1-postgres-1", "psql", "-X", "-v", "ON_ERROR_STOP=1",
         "-U", "hoistra_hw1", "-d", "hoistra_hw1", "-At"],
        input=sql, text=True, capture_output=True,
    )
    if result.returncode:
        # SQL may contain owner-provided source rows, so do not echo stderr.
        # A root-only diagnostic still lets an operator repair the importer
        # without leaking workbook content into ordinary command output.
        diagnostic = Path("/tmp/vasiliki-workbook-psql-error.log")
        diagnostic.write_text(result.stderr)
        diagnostic.chmod(0o600)
        raise RuntimeError("Workbook database operation failed; inspect the VPS privately")
    return result.stdout.strip()


def sql_literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def existing_counts(all_rows: list[SourceRow], workbook_sha: str) -> dict[str, int]:
    record_ids = [record_id(workbook_sha, row) for row in all_rows]
    payload = json.dumps(record_ids)
    query = """
SELECT json_build_object(
  'records', (SELECT count(*) FROM hw1_compliance.records WHERE id IN
      (SELECT value FROM json_array_elements_text(%s::json))),
  'sources', (SELECT count(*) FROM hw1_compliance.sources WHERE id IN
      (SELECT value FROM json_array_elements_text(%s::json))));
""" % (sql_literal(payload), sql_literal(json.dumps([source_id(workbook_sha, row) for row in all_rows])))
    raw = run_psql(query)
    result = json.loads(raw or "{}")
    return {"records": int(result.get("records", 0)), "sources": int(result.get("sources", 0))}


def existing_projection_count(projections: list[dict[str, Any]]) -> int:
    identifiers = [row["record_id"] for row in projections]
    query = "SELECT count(*) FROM hw1_access.read_records WHERE record_id IN " \
            "(SELECT value FROM json_array_elements_text(%s::json));" % sql_literal(json.dumps(identifiers))
    return int(run_psql(query) or 0)


def batches(values: list[dict[str, Any]], size: int = 250) -> Iterable[list[dict[str, Any]]]:
    # Smaller JSON statements keep PostgreSQL memory bounded while the outer
    # transaction still guarantees an all-or-nothing workbook commit.
    for start in range(0, len(values), size):
        yield values[start:start + size]


def populate_statements(table: str, values: list[dict[str, Any]], *, conflict: bool = False) -> list[str]:
    suffix = " ON CONFLICT DO NOTHING" if conflict else ""
    if table in {"hw1_compliance.sources", "hw1_compliance.records"}:
        # Exclude created_at so PostgreSQL applies its trusted server timestamp.
        prefix = f"INSERT INTO {table} (id, kind, org, buildings, data) "
        select = "SELECT id, kind, org, buildings, data FROM "
    else:
        prefix = f"INSERT INTO {table} "
        select = "SELECT * FROM "
    return [
        prefix + select + f"json_populate_recordset(NULL::{table}, "
        + sql_literal(json.dumps(batch)) + f"::json){suffix};"
        for batch in batches(values)
    ]


def build_manifest(path: Path, target_building: str) -> tuple[dict[str, Any], dict[str, list[SourceRow]]]:
    workbook_sha = file_digest(path)
    rows, errors = read_rows(path)
    errors.extend(validate_references(rows))
    all_rows = flatten(rows)
    source_buildings = {row.fields.get("building_id") for row in all_rows if "building_id" in row.fields}
    if source_buildings - {SOURCE_BUILDING_ID}:
        errors.append(f"unexpected source building identifiers: {sorted(source_buildings)}")
    counts = {sheet: len(rows.get(sheet, [])) for sheet in SHEET_HEADERS}
    projections = sum(len(read_projection(row, workbook_sha, target_building)) for row in all_rows)
    manifest = {
        "action": "complete_workbook_ingestion",
        "workbook_name": path.name,
        "workbook_sha256": workbook_sha,
        "organization_id": ORGANIZATION_ID,
        "source_building_id": SOURCE_BUILDING_ID,
        "target_building_id": target_building,
        "sheet_counts": counts,
        "total_source_rows": len(all_rows),
        "total_immutable_sources": len(all_rows),
        "total_typed_records": len(all_rows),
        "total_agent_read_projections": projections,
        "sensitive_rows_withheld_from_agent_reads": sum(len(rows.get(sheet, [])) for sheet in SENSITIVE_SHEETS),
        "validation_errors": errors,
        "status": "valid" if not errors else "invalid",
        "mapping_version": "vasiliki-complete-workbook-v1",
    }
    manifest["preview_hash"] = digest(manifest)
    return manifest, rows


def persist_preview(manifest: dict[str, Any]) -> str:
    proposal_id = "workbook-proposal-" + manifest["preview_hash"][:32]
    data = {"status": "previewed", "preview": manifest, "preview_hash": manifest["preview_hash"],
            "actor": "hw1-portfolio", "role": "portfolio_manager"}
    sql = """
BEGIN;
INSERT INTO hw1_compliance.proposals (id, kind, org, buildings, data)
VALUES (%s, 'workbook_ingestion', %s, ARRAY[%s]::text[], %s::jsonb)
ON CONFLICT (id) DO NOTHING;
COMMIT;
""" % (sql_literal(proposal_id), sql_literal(ORGANIZATION_ID),
         sql_literal(manifest["target_building_id"]), sql_literal(json.dumps(data)))
    run_psql(sql)
    return proposal_id


def confirm(manifest: dict[str, Any], rows: dict[str, list[SourceRow]], expected_hash: str, reason: str) -> dict[str, Any]:
    if manifest["status"] != "valid":
        raise ValueError("Cannot confirm a workbook with validation errors")
    if expected_hash != manifest["preview_hash"]:
        raise ValueError("Preview hash does not match the unchanged workbook batch")
    if len(reason.strip()) < 8:
        raise ValueError("A specific confirmation reason is required")

    workbook_sha = manifest["workbook_sha256"]
    target = manifest["target_building_id"]
    all_rows = flatten(rows)
    existing = existing_counts(all_rows, workbook_sha)
    if existing["records"] or existing["sources"]:
        raise ValueError("This workbook batch already has stored rows; no partial replay is allowed")

    sources = [{"id": source_id(workbook_sha, row), "kind": "workbook_" + row.sheet,
                "org": ORGANIZATION_ID, "buildings": [target], "data": source_payload(workbook_sha, row)}
               for row in all_rows]
    records = [{"id": record_id(workbook_sha, row), "kind": "workbook_" + row.sheet,
                "org": ORGANIZATION_ID, "buildings": [target], "data": record_payload(workbook_sha, row)}
               for row in all_rows]
    projections = [projection for row in all_rows for projection in read_projection(row, workbook_sha, target)]
    if existing_projection_count(projections):
        raise ValueError("Agent-readable record identifiers already exist; no partial or ambiguous projection is allowed")
    proposal_id = "workbook-proposal-" + manifest["preview_hash"][:32]
    history_id = "workbook-history-" + manifest["preview_hash"][:32]
    result = {"outcome": "accepted", "proposal_id": proposal_id, "workbook_sha256": workbook_sha,
              "source_rows": len(sources), "records": len(records), "read_projections": len(projections),
              "reason": reason.strip()}
    final_data = {"status": "confirmed", "preview": manifest, "preview_hash": expected_hash,
                  "actor": "hw1-portfolio", "role": "portfolio_manager",
                  "decision_reason": reason.strip(), "result": result}
    statements = ["BEGIN;", """DO $$ BEGIN
  IF NOT EXISTS (
    SELECT 1 FROM hw1_compliance.proposals
    WHERE id = %s AND kind = 'workbook_ingestion'
      AND data->>'status' = 'previewed' AND data->>'preview_hash' = %s
  ) THEN RAISE EXCEPTION 'Workbook preview is missing, changed, or already confirmed'; END IF;
END $$;""" % (sql_literal(proposal_id), sql_literal(expected_hash))]
    statements.extend(populate_statements("hw1_compliance.sources", sources))
    statements.extend(populate_statements("hw1_compliance.records", records))
    statements.extend(populate_statements("hw1_access.read_records", projections))
    statements.append("""INSERT INTO hw1_compliance.history (id, kind, org, buildings, data)
VALUES (%s, 'workbook_ingestion', %s, ARRAY[%s]::text[], %s::jsonb);
UPDATE hw1_compliance.proposals SET data = %s::jsonb WHERE id = %s;
COMMIT;""" % (
        sql_literal(history_id),
        sql_literal(ORGANIZATION_ID), sql_literal(target), sql_literal(json.dumps(result)),
        sql_literal(json.dumps(final_data)), sql_literal(proposal_id),
    ))
    run_psql("\n".join(statements))
    return result


def parse_args(argv: Iterable[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("mode", choices=("preview", "confirm"))
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--building-id", required=True)
    parser.add_argument("--preview-hash")
    parser.add_argument("--reason")
    parser.add_argument("--output", type=Path)
    return parser.parse_args(argv)


def main(argv: Iterable[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    if not args.workbook.is_file():
        raise FileNotFoundError(f"Workbook not found: {args.workbook}")
    manifest, rows = build_manifest(args.workbook, args.building_id)
    if args.mode == "preview":
        manifest["existing_batch_rows"] = existing_counts(flatten(rows), manifest["workbook_sha256"])
        proposal_id = persist_preview(manifest) if manifest["status"] == "valid" else None
        output = {"proposal_id": proposal_id, **manifest}
    else:
        if not args.preview_hash or not args.reason:
            raise ValueError("confirm requires --preview-hash and --reason")
        output = confirm(manifest, rows, args.preview_hash, args.reason)
    rendered = json.dumps(output, indent=2, sort_keys=True)
    print(rendered)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
